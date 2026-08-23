from core.forecasting.valuation import dcf_sensitivity


def build_sensitivity_matrix(
    free_cash_flows,
    discount_rates,
    terminal_growth_rates,
):
    """
    Build a structured DCF sensitivity matrix.

    Rows    = discount rates
    Columns = terminal growth rates
    """

    raw_sensitivity = dcf_sensitivity(
        free_cash_flows=free_cash_flows,
        discount_rates=discount_rates,
        terminal_growth_rates=terminal_growth_rates,
    )

    matrix = []

    for discount_rate in discount_rates:

        row = {
            "discount_rate": discount_rate,
            "values": [],
        }

        for growth_rate in terminal_growth_rates:

            enterprise_value = raw_sensitivity[
                discount_rate
            ][growth_rate]

            row["values"].append(
                {
                    "terminal_growth_rate": growth_rate,
                    "enterprise_value": enterprise_value,
                }
            )

        matrix.append(row)

    return matrix
from openpyxl import Workbook


def write_sensitivity_matrix(
    workbook,
    matrix,
):
    """
    Write a DCF sensitivity matrix into the
    Sensitivity worksheet.
    """

    sheet = workbook["Sensitivity"]

    sheet["A1"] = "DCF Sensitivity Analysis"

    sheet["A3"] = "WACC / Terminal Growth"

    if not matrix:
        return sheet

    growth_rates = [
        item["terminal_growth_rate"]
        for item in matrix[0]["values"]
    ]

    for column, growth_rate in enumerate(
        growth_rates,
        start=2,
    ):
        sheet.cell(
            row=3,
            column=column,
            value=growth_rate,
        )

    for row, sensitivity_row in enumerate(
        matrix,
        start=4,
    ):
        discount_rate = sensitivity_row[
            "discount_rate"
        ]

        sheet.cell(
            row=row,
            column=1,
            value=discount_rate,
        )

        for column, value_data in enumerate(
            sensitivity_row["values"],
            start=2,
        ):
            sheet.cell(
                row=row,
                column=column,
                value=round(
                    value_data["enterprise_value"],
                    2,
                ),
            )

    return sheet