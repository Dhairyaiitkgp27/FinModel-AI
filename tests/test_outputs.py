"""Tests for Excel and PDF report generation."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.analysis import AnalysisAssumptions, AnalysisBundle, run_full_analysis
from core.outputs import build_excel_report, build_pdf_report

SAMPLE_DIR = Path(__file__).resolve().parents[1] / "data" / "sample"


@pytest.fixture(scope="module")
def bundle() -> AnalysisBundle:
    return run_full_analysis("AAPL", provider="local", sample_dir=SAMPLE_DIR)


# --------------------------------------------------------------------------- #
# Excel                                                                        #
# --------------------------------------------------------------------------- #
def test_excel_report_has_expected_sheets(bundle: AnalysisBundle, tmp_path: Path):
    path = build_excel_report(bundle, tmp_path / "report.xlsx")
    assert path.exists()
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Assumptions", "Ratios", "Forecast", "Valuation", "Monte Carlo"]


def test_excel_summary_contains_company_and_dcf(bundle: AnalysisBundle, tmp_path: Path):
    path = build_excel_report(bundle, tmp_path / "report.xlsx")
    wb = load_workbook(path)
    summary = wb["Summary"]
    assert summary["A1"].value == "Apple Inc. (AAPL)"
    dcf_value = None
    for row in summary.iter_rows(values_only=True):
        if row and row[0] == "DCF value / share":
            dcf_value = row[1]
    assert dcf_value == pytest.approx(bundle.dcf.implied_share_price, abs=0.01)


def test_excel_forecast_sheet_has_line_items(bundle: AnalysisBundle, tmp_path: Path):
    path = build_excel_report(bundle, tmp_path / "report.xlsx")
    wb = load_workbook(path)
    labels = {row[0] for row in wb["Forecast"].iter_rows(values_only=True) if row and row[0]}
    assert {"Revenue", "EBITDA", "Unlevered FCFF", "Total assets"} <= labels


def test_excel_monte_carlo_sheet_has_percentiles(bundle: AnalysisBundle, tmp_path: Path):
    path = build_excel_report(bundle, tmp_path / "report.xlsx")
    wb = load_workbook(path)
    labels = {row[0] for row in wb["Monte Carlo"].iter_rows(values_only=True) if row and row[0]}
    assert "P10" in labels and "P90" in labels


# --------------------------------------------------------------------------- #
# PDF                                                                          #
# --------------------------------------------------------------------------- #
def test_pdf_report_is_valid_file(bundle: AnalysisBundle, tmp_path: Path):
    path = build_pdf_report(bundle, tmp_path / "report.pdf")
    assert path.exists()
    data = path.read_bytes()
    assert data[:5] == b"%PDF-"
    assert len(data) > 1000  # a non-trivial document


# --------------------------------------------------------------------------- #
# Graceful degradation                                                         #
# --------------------------------------------------------------------------- #
def _minimal_bundle() -> AnalysisBundle:
    from core.data.dataset import CompanyDataset
    from core.data.validation.checks import ValidationReport
    from core.models import CompanyProfile, FinancialStatements, PriceHistory
    from core.models.analysis import RatioAnalysis

    return AnalysisBundle(
        dataset=CompanyDataset(
            profile=CompanyProfile(ticker="TST", name="Test Co"),
            financials=FinancialStatements(ticker="TST"),
            prices=PriceHistory(ticker="TST"),
            source="test",
        ),
        assumptions=AnalysisAssumptions(),
        validation=ValidationReport(ticker="TST"),
        ratios=RatioAnalysis(ticker="TST"),
    )


def test_reports_handle_missing_sections(tmp_path: Path):
    # A bundle with no valuations should still produce both files without error.
    bundle = _minimal_bundle()
    xlsx = build_excel_report(bundle, tmp_path / "empty.xlsx")
    pdf = build_pdf_report(bundle, tmp_path / "empty.pdf")
    assert xlsx.exists()
    assert pdf.read_bytes()[:5] == b"%PDF-"
